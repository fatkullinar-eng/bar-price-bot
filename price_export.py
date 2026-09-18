# ============================================================
# УПАКОВЩИК ОТЧЁТОВ — price_export.py
# Собирает ПОЛНЫЙ список позиций из истории (поиск — как у /find,
# но без ограничения в 15 строк) и записывает его в файл Excel.
# В чат бот потом отправляет этот файл.
#
# Этот файл ничего не отправляет в Telegram — «отправкой»
# занимается bot.py. Здесь только поиск и запись в Excel.
# ============================================================

import re                        # шаблоны поиска и чистка имени файла
from datetime import date        # сегодняшняя дата — для имени отчёта
from pathlib import Path         # работа с путями к файлам

from openpyxl import Workbook    # Excel-упаковщик (ставили на Этапе 2)
from openpyxl.styles import Font # стиль текста в ячейках (жирная шапка)

from price_report import (       # готовые помощники аналитика:
    load_history,                #   прочитать все CSV истории
    row_matches,                 #   подходит ли строка под запрос
    parse_query,                 #   «хеллес до 200» -> слова + потолок
    to_float,                    #   «137.75» -> число
    clean,                       #   убрать переносы строк из названия
    is_soft_drink,               #   не-пиво? (лимонады, вода, квас…)
    fmt_price,                   #   число для чата: 137.75 -> «137,75»
    who,                         #   подпись «пивоварня — поставщик»
)

# Заголовки колонок будущего Excel-файла
HEADERS = [
    "Поставщик", "Дата", "Название", "Пивоварня", "Стиль",
    "Тара", "Объём, л", "Цена, руб", "Цена за литр, руб",
]

# Ширина колонок (чтобы всё было видно без «растягивания»)
WIDTHS = [15, 12, 45, 20, 20, 14, 10, 10, 14]


def export_report(query: str, history_dir, reports_dir):
    """Ищет по истории (правила — как у /find) и упаковывает ВСЁ
    найденное в Excel-файл с двумя листами: кеги и банки.

    Возвращает (путь к файлу, короткий текст для чата),
    а если ничего не нашлось — (None, текст-объяснение)."""
    words, cap, brewery_only = parse_query(query)
    if not words:
        return None, (
            "Напишите, что искать. Пример: /report хеллес\n"
            "Соберу ВСЕ найденные позиции (без ограничения в 15 строк) "
            "и пришлю файлом Excel."
        )

    rows = load_history(history_dir)
    if not rows:
        return None, "В истории пока пусто — начните с того, что пришлите прайс."

    # Отбираем строки: совпал запрос; цена влезла в потолок (если указан).
    # Те же правила, что у /find — чтобы «найти» и «отчёт» не спорили.
    matched = []
    for r in rows:
        if not row_matches(r, words, brewery_only):
            continue
        per_liter = to_float(r["цена за литр, руб"])
        price = to_float(r["цена, руб"])
        if cap is not None:
            limit = per_liter if r["большая тара (кег/пэт)"] == "да" else price
            if limit is None or limit > cap:
                continue
        matched.append(r)

    if not matched:
        return None, (
            f"По запросу «{' '.join(words)}» ничего не нашёл у всех "
            "поставщиков в истории. Попробуйте другое слово или уберите "
            "цену-потолок."
        )

    # Сортировка: кеги — по цене за литр, банки — по цене за штуку.
    # Позиции без цены честно ставим в конец списка, не выбрасываем.
    def sort_key(r, field):
        v = to_float(r[field])
        return (v is None, v if v is not None else 0)

    kegs = sorted(
        (r for r in matched if r["большая тара (кег/пэт)"] == "да"),
        key=lambda r: sort_key(r, "цена за литр, руб"),
    )
    cans = sorted(
        (r for r in matched if r["большая тара (кег/пэт)"] == "нет"),
        key=lambda r: sort_key(r, "цена, руб"),
    )

    # --- Собираем Excel-книгу: два листа ---
    wb = Workbook()
    ws_kegs = wb.active                       # первый лист создаётся сам
    ws_kegs.title = "Кеги и пэты (за литр)"
    ws_cans = wb.create_sheet("Банки и бутылки (за штуку)")
    _fill_sheet(ws_kegs, kegs)
    _fill_sheet(ws_cans, cans)

    # --- Имя файла: «2026-09-14 — отчёт хеллес — позиции.xlsx» ---
    today = date.today().isoformat()
    # Из запроса делаем «безопасное» имя: убираем символы, которые
    # Windows запрещает в именах файлов (слэши, двоеточия и т.п.)
    slug = re.sub(r'[\\/:*?"<>|]+', " ", " ".join(words)).strip()
    slug = re.sub(r"\s+", " ", slug)[:40]     # без двойных пробелов, до 40 знаков
    path = Path(reports_dir) / f"{today} — отчёт {slug} — позиции.xlsx"
    wb.save(path)

    # --- Короткий текст для чата (выжимка) ---
    sup_counts: dict = {}
    for r in matched:
        sup_counts[r["поставщик"]] = sup_counts.get(r["поставщик"], 0) + 1
    by_sup = ", ".join(f"{s} — {n}" for s, n in sorted(sup_counts.items()))
    chat_text = (
        f"🔎 Нашёл: {len(kegs)} кегов и пэт, {len(cans)} банок и бутылок.\n"
        f"У поставщиков: {by_sup}\n\n"
        f"Полный список — в файле Excel (два листа: кеги и банки)."
    )
    return path, chat_text


def _fill_sheet(ws, rows: list) -> None:
    """Записывает строки истории на один лист Excel: шапка + данные."""
    # Шапка — жирным
    for col, title in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True)
    # Ширина колонок — чтобы читалось без растягивания
    for col, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width
    # Первая строка закреплена — при прокрутке шапка остаётся видимой
    ws.freeze_panes = "A2"

    for i, r in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=r["поставщик"])
        ws.cell(row=i, column=2, value=r["дата"])
        ws.cell(row=i, column=3, value=clean(r["название"]))
        ws.cell(row=i, column=4, value=clean(r["пивоварня"]))
        ws.cell(row=i, column=5, value=clean(r["стиль"]))
        ws.cell(
            row=i, column=6,
            value="кег/пэт" if r["большая тара (кег/пэт)"] == "да" else "банка/бутылка",
        )
        ws.cell(row=i, column=7, value=to_float(r["объём штуки, л"]))
        ws.cell(row=i, column=8, value=to_float(r["цена, руб"]))
        ws.cell(row=i, column=9, value=to_float(r["цена за литр, руб"]))


def export_sorted_report(history_dir, reports_dir, supplier: str = ""):
    """«Сортировка» (/sort, слово «сортировка»): пиво из истории,
    выстроенное по возрастанию цены — БЕЗ разделения на стили.
      • розлив (кеги и пэты) — по цене за литр;
      • банки и бутылки — по цене за штуку.

    supplier: пусто — все поставщики вместе; иначе — только один
    поставщик (имя можно писать неточно: «бирдилер» найдёт «Бирдилер»).

    Не-пиво (лимонады, вода, квас…) исключаем — то же правило, что в топах.

    Возвращает (путь к файлу Excel, текст для чата с выжимкой),
    а если истории нет / поставщик не найден — (None, текст-объяснение)."""
    rows = load_history(history_dir)
    if not rows:
        return None, "В истории пока пусто — начните с того, что пришлите прайс."

    # --- Фильтр по поставщику (если назвали конкретного) ---
    wanted = (supplier or "").strip()
    title = "по всем поставщикам"
    if wanted:
        low = wanted.casefold()
        # Сначала точное совпадение, потом «часть слова» (бирдилер -> Бирдилер)
        exact = [r for r in rows if r["поставщик"].casefold() == low]
        if exact:
            rows = exact
        else:
            matches = sorted({
                r["поставщик"] for r in rows
                if low in r["поставщик"].casefold()
            })
            if len(matches) == 1:
                rows = [r for r in rows if r["поставщик"] == matches[0]]
            elif len(matches) > 1:
                return None, (
                    "По этому слову несколько поставщиков:\n  "
                    + "\n  ".join(matches)
                    + "\nУточните, пожалуйста."
                )
            else:
                all_sups = sorted({r["поставщик"] for r in rows})
                return None, (
                    f"Поставщика «{wanted}» в истории нет.\n"
                    "Кто есть:\n  " + "\n  ".join(all_sups)
                )
        sup_name = rows[0]["поставщик"]       # «правильное» написание
        title = f"«{sup_name}»"

    # Только пиво/сидр/мёд — как в топах дайджеста
    beer = [r for r in rows if not is_soft_drink(r)]
    if not beer:
        return None, (
            "В истории только не-пиво (лимонады, вода и т.п.) — "
            "сортировать нечего."
        )

    # Сортировка по возрастанию: кеги — за литр, банки — за штуку.
    # Позиции без цены честно ставим в конец, не выбрасываем.
    def sort_key(r, field):
        v = to_float(r[field])
        return (v is None, v if v is not None else 0)

    kegs = sorted(
        (r for r in beer if r["большая тара (кег/пэт)"] == "да"),
        key=lambda r: sort_key(r, "цена за литр, руб"),
    )
    cans = sorted(
        (r for r in beer if r["большая тара (кег/пэт)"] == "нет"),
        key=lambda r: sort_key(r, "цена, руб"),
    )

    # --- Собираем Excel-книгу: два листа ---
    wb = Workbook()
    ws_kegs = wb.active                       # первый лист создаётся сам
    ws_kegs.title = "Розлив (за литр)"
    ws_cans = wb.create_sheet("Банки и бутылки (за штуку)")
    _fill_sheet(ws_kegs, kegs)
    _fill_sheet(ws_cans, cans)

    today = date.today().isoformat()
    # Имя файла: с поставщиком, если сортировали конкретного
    name_part = f" {rows[0]['поставщик']}" if wanted else " по цене"
    path = Path(reports_dir) / f"{today} — сортировка{name_part} — позиции.xlsx"
    wb.save(path)

    # --- Выжимка для чата: первые 10 самых дешёвых в каждой группе ---
    keg_lines = []
    for r in kegs[:10]:
        v = to_float(r["цена за литр, руб"])
        if v is None:
            break   # без цены всегда в конце — значит, дешевле не осталось
        keg_lines.append(
            f"  • {fmt_price(v)} ₽/л — {clean(r['название'])} ({who(r)})"
        )
    can_lines = []
    for r in cans[:10]:
        v = to_float(r["цена, руб"])
        if v is None:
            break
        can_lines.append(
            f"  • {fmt_price(v)} ₽ — {clean(r['название'])} ({who(r)})"
        )

    chat_text = (
        f"📋 Сортировка по цене {title} (только пиво, без разделения на стили).\n"
        f"Розлив: {len(kegs)} позиций. Самые дешёвые:\n"
        + "\n".join(keg_lines)
        + f"\n\nБанки и бутылки: {len(cans)} позиций. Самые дешёвые:\n"
        + "\n".join(can_lines)
        + "\n\nПолный список — в файле Excel: два листа, от дешёвых к дорогим."
    )
    return path, chat_text