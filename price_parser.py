# ============================================================
# ЧИТАТЕЛЬ ПРАЙСОВ — price_parser.py
# Открывает прайс (Excel или PDF), находит таблицы с товарами
# и вытаскивает из каждой строки: пивоварню, название, стиль,
# тару, объём, цену. Считает цену за литр.
# Кривые строки не ломают чтение — они пропускаются
# и честно пересчитываются.
#
# Правило бара: кеги/пэт показываем ЗА ЛИТР,
#              банки/бутылки — ЗА ШТУКУ.
# ============================================================

import csv                    # простейший формат таблиц: сохраняем историю
import re                     # «шаблоны поиска» в тексте (регулярные выражения)
import unicodedata           # приводит «двойные» невидимые буквы к обычным
from datetime import date
from pathlib import Path

from openpyxl import load_workbook   # читалка Excel
import pdfplumber                    # читалка PDF: находит таблицы на страницах

from config import SUPPLIER_ALIASES            # справочник имён поставщиков (config.py)
from config import PET_PER_LITER_SUPPLIERS     # у кого цена ПЭТ написана ЗА ЛИТР

# ---------- Шаблоны поиска ----------
# Регулярное выражение — это трафарет для поиска в тексте.
# Первый трафарет ловит число перед "л": "кег 20л" -> 20
RE_LITERS = re.compile(r"(\d+(?:[.,]\d+)?)\s*л", re.IGNORECASE)
# Второй трафарет ловит упаковку вида "кор. 12х0,45л":
# -> 12 банок в коробке, каждая по 0.45 литра
RE_PACK = re.compile(r"(\d+)\s*[хx*]\s*(\d+(?:[.,]\d+)?)\s*л", re.IGNORECASE)
# Третий трафарет ловит объём в САМОМ НАЗВАНИИ позиции, рядом со словом
# тары: «Бутылка 0,33», «Банка 0,45», «ПЭТ 1,5», «Бут 0,44».
# Слово тары обязательно — иначе «пиво 1516» можно принять за 1516 литров.
RE_NAME_VOL = re.compile(
    r"(?:бутыл\w*|банк\w*|пэт|бут\b|ж/б|стекл\w*)\W{0,3}(\d+(?:[.,]\d+)?)",
    re.IGNORECASE,
)
# Четвёртый трафарет — «двойная цена» в ОДНОЙ ячейке: «260р / 5200р».
# Так оформлен прайс Black Cat (Google-таблица): слева цена за литр,
# справа — за кег целиком (или банка / коробка). Ловим ПАРЫ чисел
# вокруг дроби; после первого числа обязательно «р»/«руб» — иначе
# в пару превратилось бы «коробка 12 / 20 шт».
RE_DUAL_PAIR = re.compile(
    r"(\d+(?:[.,]\d+)?)\s*(?:р|руб\.?)\s*/\s*(\d+(?:[.,]\d+)?)",
    re.IGNORECASE,
)
# Пятый трафарет — объём банки внутри ячейки с ценой: «0,45 ж/б».
# Число стоит ПЕРЕД словом тары, поэтому прежние трафареты его не видят.
RE_CAN_VOL = re.compile(
    r"(\d+(?:[.,]\d+)?)\s*(?:ж/?б|банк\w*|бут\w*|стекл\w*)",
    re.IGNORECASE,
)

# Листы, которые НЕ читаем: "Заказано" — черновик для оформления заказа.
# Листы «Прочее» теперь читаем: там бывают лимонады и квас (в данные — да,
# в топы — нет: их оттуда удерживает фильтр «не-пива»).
SKIP_SHEETS = ("заказано",)

# Слова-пароли «не-пива» для топов: газировки, соки, чаи, энергетики.
# Из ДАННЫХ и ПОИСКА они не исчезают — убираем только из списка «самое дешёвое».
# Список пополняется: встретилась новая марка газировки — добавили слово.
SOFT_WORDS = (
    "лимонад", "lemonade", "квас",       # договорились сразу: квас и лимонады
    "тархун", "мохито", "шорли",         # газировки-коктейли
    "чай", "tea",                        # чаи
    "энергетик", "энерг",                # энергетики
    "schweppes", "dr pepper", "fanta",   # известные марки газировок
    "chupa chups", "cola",               # и ещё две на всякий случай
    # Соки — БЕЗ слова «сок»! Оно коварно: прячется внутри «выСОКой
    # плотности» и «МедоНОСОК» (а это мёд, ему место в топах!),
    # а «В собственном соку» и «Alpaca Juice» — это вообще пиво.
    # Настоящие соки ловим по словам «напитки/напиток»: у пивомира
    # соки Vinut лежат на листе «Безалкогольные напитки».
    "напитки", "напиток",
    # Вода — тоже не пиво. ВАЖНО: короткое слово «вода» сюда нельзя —
    # оно прячется внутри «пивЗАВОДА» и «завода», и из топов исчезло бы
    # всё от «Пензенского пивзавода». Ловим воду по её приметам:
    "газированная", "негазированная",    # «Хрустальная газированная…»
    "хрустальная",                       # вода «Хрустальная» из прайсов ЗТ
    "hop water",                         # вода с хмелем из прайса Price_Craft
)

# Строку-шапку (с названиями колонок) ищем не глубже этого числа строк:
MAX_HEADER_SCAN = 12


# ---------- Мелкие помощники ----------

def to_text(value) -> str:
    """Превращает ячейку в текст. Пустая ячейка -> пустая строка."""
    return "" if value is None else str(value).strip()


def to_number(value):
    """Превращает ячейку с ценой в число. Не получилось -> None."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    text = text.replace(" ", "").replace("\xa0", "")   # убираем пробелы: "1 200"
    text = text.replace(",", ".")                        # запятая -> точка
    digits = re.sub(r"[^0-9.]", "", text)                 # оставляем только цифры и точку
    if not digits or digits == ".":
        return None
    try:
        return float(digits)
    except ValueError:
        return None


# Страны, которые поставщики пишут рядом с пивоварней в заголовках разделов
# («4brewers Россия», «Achel Бельгия») — в поле «пивоварня» стране не место.
COUNTRIES = (
    "россия", "бельгия", "ирландия", "германия", "чехия", "англия",
    "великобритания", "сша", "италия", "франция", "польша", "литва",
    "латвия", "эстония", "дания", "австрия", "испания", "португалия",
    "нидерланды", "финляндия", "швеция", "япония", "китай",
    "мексика", "тайвань", "соединенное", "королевство",
)


def clean_brewery(text: str) -> str:
    """Чистит заголовок раздела: «Achel Бельгия» -> «Achel»."""
    words = [w for w in text.split() if w.lower().strip(".,()") not in COUNTRIES]
    return " ".join(words)


def guess_supplier(file_name: str) -> str:
    """Угадывает название поставщика по имени файла.
    «2026-09-14 — Прайс Бирдилер 14.09.2026» -> «Бирдилер».
    Ничего не осталось (в имени нет названия) -> пустая строка —
    тогда бот спросит поставщика у человека прямо в чате."""
    text = Path(file_name).stem            # имя без расширения
    # Telegram иногда сохраняет «й» как ДВА невидимых символа —
    # глазом не видно, но наш фильтр «прайс» её не узнаёт.
    # Приводим буквы к обычному виду (NFC — «нормальная форма»).
    text = unicodedata.normalize("NFC", text)
    # Сначала — справочник (config.py): знакомое слово в имени файла
    # сразу даёт чистое имя поставщика, без угадываний.
    low = text.lower()
    for fragment, supplier in SUPPLIER_ALIASES.items():
        if fragment.lower() in low:
            return supplier
    text = re.sub(r"\d{4}[-.]\d{2}[-.]\d{2}", "", text)             # дата ГГГГ-ММ-ДД
    text = re.sub(r"\d{2}[.\-/_]\d{2}[.\-/_]\d{2,4}", "", text)     # дата ДД.ММ.ГГГГ (и с подчёркиваниями)
    text = re.sub(r"(?i)\bот\s+\d{1,2}[.\-/_]\d{1,2}([.\-/_]\d{2,4})?", "", text)   # «от 14.09» после даты
    text = re.sub(r"(?i)прайс[- _]?(лист)?", "", text)              # слово «прайс»
    text = re.sub(r"(?i)[-_ ]+price[_ ]*$", "", text)               # английское «price» в КОНЦЕ имени
    text = re.sub(r"\s*\(\d+\)\s*$", "", text)                      # «(1)» в конце — метка Telegram о повторном скачивании
    text = text.strip(" -_—")                                        # лишние чёрточки по краям
    return text


def is_pet_per_liter_supplier(supplier: str) -> bool:
    """True, если поставщик есть в справочнике «цена ПЭТ за литр»
    (PET_PER_LITER_SUPPLIERS в config.py). У таких поставщиков
    в колонке «Цена» для ПЭТ-тары стоит цена ЗА ЛИТР, а не за тару
    целиком — бот это учитывает при подсчёте."""
    low = (supplier or "").lower()
    return any(frag.lower() in low for frag in PET_PER_LITER_SUPPLIERS)


# ---------- Поиск шапки таблицы ----------

def find_header_in_rows(rows):
    """Ищет строку-шапку среди первых строк таблицы (там написано
    «Название», «Цена»...). Возвращает (номер строки, карта колонок)
    или (None, None). Работает и с листами Excel, и с таблицами из PDF."""
    for row_num, row in enumerate(rows[:MAX_HEADER_SCAN], start=1):
        texts = [to_text(c).lower() for c in row]
        # Шапка — строка, где есть и «название»-подобное, и «цена»-подобное
        has_name = any(
            ("название" in t) or ("наименование" in t) or ("номенклатура" in t)
            for t in texts
        )
        has_price = any("цена" in t and "упак" not in t for t in texts)
        if not (has_name and has_price):
            continue

        cols = {}
        # Название: предпочитаем «Название», потом «Наименование», потом «Номенклатура»
        for hint in ("название", "наименование", "номенклатура"):
            for idx, t in enumerate(texts):
                if hint in t and "name" not in cols:
                    cols["name"] = idx
                    break
        for idx, t in enumerate(texts):
            if "пивоварня" in t or "производитель" in t or "сидродельня" in t:
                cols["brewery"] = idx
            elif "стиль" in t or "cтиль" in t:
                # «cтиль» с ЛАТИНСКОЙ «c» — так написано в прайсе Black Cat:
                # глазом не отличить от русской «с», а поиск — различает.
                cols["style"] = idx
            elif ("объём" in t or "объем" in t) and "volume" not in cols:
                cols["volume"] = idx        # отдельная колонка «Объём» (0.5 / 20 / 30)
            elif t.startswith("ед") and "unit" not in cols:
                cols["unit"] = idx          # «ед. изм.» — кег 20л / ж/б / бут
            elif "упаковка" in t or "тара" in t:
                cols["pack"] = idx
            elif ("цена" in t and "литр" in t and "кег" in t
                  and "price_dual" not in cols):
                # «Цена литр / кег (20л)» — ДВЕ цены в одной ячейке:
                # слева за литр, справа за кег целиком (прайс Black Cat).
                # Объём кега написан в самой шапке: «(20л)» -> 20 литров.
                cols["price_dual"] = idx
                m = RE_LITERS.search(t)
                if m:
                    v = float(m.group(1).replace(",", "."))
                    if v >= 2:                      # 20 л, 30 л — точно большая тара
                        cols["dual_keg_liters"] = v
            elif ("цена" in t and "банк" in t and "короб" in t
                  and "price_dual_can" not in cols):
                # «Цена банка 0,33 - 0,45 / коробка (12 шт)» — слева цена
                # банки, справа коробки. Коробку не считаем — берём банками.
                cols["price_dual_can"] = idx
            elif "цена" in t and "лит" in t and "price_liter" not in cols:
                cols["price_liter"] = idx   # «Цена литр»/«ЦЕНА ЗА 1 лит.» — цена уже ЗА ЛИТР
            elif "цена" in t and "кег" in t and "price_keg" not in cols:
                cols["price_keg"] = idx     # «Цена Кега» — за кег целиком
            elif "цена" in t and "упак" not in t and "price" not in cols:
                cols["price"] = idx
        # Лист считается таблицей с товарами, если есть название
        # и ХОТЯ БЫ какая-нибудь ценовая колонка
        if "name" in cols and (
            "price" in cols or "price_liter" in cols or "price_keg" in cols
            or "price_dual" in cols or "price_dual_can" in cols
        ):
            return row_num, cols
    return None, None


def find_header(ws):
    """Ищет в листе Excel строку-шапку — заглядывает в первые строки."""
    rows = list(ws.iter_rows(max_row=MAX_HEADER_SCAN, values_only=True))
    return find_header_in_rows(rows)


# ---------- Разбор одной строки ----------

def parse_row(row, cols, pet_per_liter=False):
    """Разбирает одну строку прайса. Возвращает позицию (словарь) или None.

    pet_per_liter — особенность поставщика (см. config.py): в его прайсе
    цена ПЭТ-тары написана ЗА ЛИТР. Для остальных поставщиков False."""
    def cell(key):
        idx = cols.get(key)
        if idx is None or idx >= len(row):
            return ""
        return to_text(row[idx])

    def num_cell(key):
        idx = cols.get(key)
        if idx is None or idx >= len(row):
            return None
        return to_number(row[idx])

    name = cell("name")
    if not name:
        return None    # нет названия — строка не товар (заголовок раздела, мусор)
    # Строка-шапка может встретиться и в СЕРЕДИНЕ таблицы (PDF иногда
    # рисует её в конце страницы) — не примем её за товар
    if name.lower() in ("название", "наименование", "номенклатура"):
        return None

    unit_text = cell("unit")     # «кег 20л», «ж/б», «бут»
    pack_text = cell("pack")     # «кор. 12х0,45л», «ПЭТ», «бут»

    liters = None    # объём ОДНОЙ штуки (или кега) в литрах

    # 1) Отдельная колонка «Объём» с числом: 0.5 / 20 / 30
    vol = num_cell("volume")
    if vol:
        liters = vol

    # 2) Объём прямо в «ед. изм.»: «кег 20л» -> 20 литров
    if liters is None:
        m = RE_LITERS.search(unit_text)
        if m:
            v = float(m.group(1).replace(",", "."))
            if v >= 2:                      # 20 л, 30 л — точно большая тара
                liters = v

    # 3) Упаковка вида «кор. 12х0,45л» -> банка 0.45 л
    if liters is None:
        m = RE_PACK.search(pack_text)
        if m:
            liters = float(m.group(2).replace(",", "."))

    # 4) На всякий случай: просто число с «л» в упаковке («бутылка 0,5л»)
    if liters is None:
        m = RE_LITERS.search(pack_text)
        if m:
            liters = float(m.group(1).replace(",", "."))

    # 5) Объём в САМОМ НАЗВАНИИ: «Балтика №0 Бутылка 0,33» -> 0.33 л.
    # Так бывает, когда в прайсе нет колонок с объёмом (наборы ЗТ).
    # Сначала ищем число рядом со словом тары («Бутылка 0,33») — это надёжно;
    # потом — просто число с «л» в названии («Вайцен 0,5л»).
    if liters is None:
        m = RE_NAME_VOL.search(name)
        if m:
            liters = float(m.group(1).replace(",", "."))
    if liters is None:
        m = RE_LITERS.search(name)
        if m:
            liters = float(m.group(1).replace(",", "."))

    # Большая тара? Пэт/кег в тексте, или объём явно большой (20–30 л)
    tare_text = (unit_text + " " + pack_text).lower()
    keg = (liters is not None and liters >= 2) or ("пэт" in tare_text) or ("кег" in tare_text)

    # ---------- Цена ----------
    # Поставщики пишут цену по-разному:
    #   «Цена»      — за банку ИЛИ за кег целиком (Бирдилер);
    #   «Цена литр» — уже ЗА ЛИТР (Price_Craft);
    #   «Цена Кега» — за кег целиком.
    price_liter = num_cell("price_liter")
    price_keg = num_cell("price_keg")
    price_plain = num_cell("price")

    price = None            # цена за штуку или за кег целиком
    price_per_liter = None  # наша главная валюта сравнения

    if price_liter:
        # Поставщик дал отдельную колонку «Цена литр» — верим ей
        price_per_liter = price_liter
        # Цена за кег/набор целиком: сперва готовые колонки поставщика
        # («Цена Кега», «Цена за набор»), и лишь если их нет — считаем сами.
        # Небольшая разница бывает из-за округления у поставщика,
        # его цифры важнее наших расчётов.
        price = (
            price_keg
            or price_plain
            or (round(price_liter * liters, 2) if liters else None)
        )
    elif keg and price_keg:
        # Есть «Цена Кега» — цена за кег целиком
        price = price_keg
        price_per_liter = round(price_keg / liters, 2) if liters else None
    elif price_plain:
        # Обычная колонка «Цена». Но внимание: иногда в «ед. изм.»
        # стоит слово «Литр» — тогда цена написана ЗА ЛИТР
        # (так делает Бирдилер в сидрах: 137 ₽/л, кег 30 л = 4110 ₽).
        if pet_per_liter and "пэт" in tare_text:
            # Прайс этого поставщика пишет цену ПЭТ ЗА ЛИТР
            # (Craft Cartel: «ПЭТ 30 л — 151 ₽», то есть 151 ₽/л,
            # а вся тара стоит 151 × 30 = 4530 ₽). Пересчитываем
            # цену тары сами. Банки читаются как обычно — за штуку.
            price_per_liter = price_plain
            price = round(price_plain * liters, 2) if liters else None
        elif "литр" in unit_text.lower():
            price_per_liter = price_plain
            price = round(price_plain * liters, 2) if liters else None
        else:
            price = price_plain
            price_per_liter = round(price_plain / liters, 2) if liters else None

    # Нет цены вообще — строка нам не товар
    if price is None or price <= 0:
        return None

    return {
        "brewery": cell("brewery"),
        "name": name,
        "style": cell("style"),
        "unit": unit_text or pack_text,   # как называется тара
        "liters": liters,                 # объём штуки (банки/кега) в литрах
        "keg": keg,                       # кег/пэт — или банка/бутылка
        "price": price,                   # цена за банку ИЛИ за кег
        "price_per_liter": price_per_liter,
    }


# ---------- Разбор строки с «двойными» ценами ----------
# Прайс Black Cat (Google-таблица пивоварни) устроен иначе, чем обычно:
# в ОДНОЙ ячейке живут две цены через дробь:
#   «260р / 5200р»                     = 260 ₽ за литр / 5200 ₽ за кег 20 л
#   «168р / 2016 (коробка 12 шт) 0,45 ж/б» = 168 ₽ банка / 2016 ₽ коробка
# Из одной строки товара получаются ДВЕ позиции: кег и/или банка.
# Если пар несколько («389р/7780р −20% 311р/6220р») — берём ПОСЛЕДНЮЮ:
# это действующая цена со скидкой. «Нет в наличии» (пар нет) — упаковка
# честно пропускается.

def parse_dual_row(row, cols):
    """Разбирает строку прайса с двойными ценами.
    Возвращает СПИСОК позиций (кег и/или банка) — может быть пустым."""
    def cell(key):
        idx = cols.get(key)
        if idx is None or idx >= len(row):
            return ""
        return to_text(row[idx])

    name = cell("name")
    if not name:
        return []
    # Строка-шапка может встретиться и внутри таблицы — не примем за товар
    if name.lower() in ("название", "наименование", "номенклатура"):
        return []
    style = cell("style")

    items = []

    # --- Кег: колонка «Цена литр / кег (20л)» ---
    idx = cols.get("price_dual")
    if idx is not None and idx < len(row):
        text = to_text(row[idx])
        pairs = RE_DUAL_PAIR.findall(text)
        if pairs:
            per_liter, keg_price = pairs[-1]
            per_liter = float(per_liter.replace(",", "."))
            keg_price = float(keg_price.replace(",", "."))
            liters = cols.get("dual_keg_liters")   # объём кега из шапки (20 л)
            if keg_price > 0 and per_liter > 0:
                items.append({
                    "brewery": "",
                    "name": name,
                    "style": style,
                    "unit": f"кег {liters:g} л" if liters else "кег",
                    "liters": liters,
                    "keg": True,
                    "price": keg_price,            # цена кега целиком
                    "price_per_liter": per_liter,   # цена за литр — от пивоварни
                })

    # --- Банка: колонка «Цена банка / коробка» ---
    idx = cols.get("price_dual_can")
    if idx is not None and idx < len(row):
        text = to_text(row[idx])
        pairs = RE_DUAL_PAIR.findall(text)
        if pairs:
            can_price, box_price = pairs[-1]
            can_price = float(can_price.replace(",", "."))
            # Объём одной банки спрятан в той же ячейке: «0,45 ж/б»
            m = RE_CAN_VOL.search(text)
            liters = float(m.group(1).replace(",", ".")) if m else None
            if can_price > 0:
                items.append({
                    "brewery": "",
                    "name": name,
                    "style": style,
                    "unit": f"банка {liters:g} л" if liters else "банка",
                    "liters": liters,
                    "keg": False,
                    "price": can_price,       # цена одной банки
                    "price_per_liter": (
                        round(can_price / liters, 2) if liters else None
                    ),
                })

    return items


# ---------- Главные функции ----------

def _is_junk_row(row):
    """Строки-заготовки для оформления заказа: в них нет товара —
    только нолики и «Итого» (так Google-таблицы хранят пустые строки
    формы заказа, их бывают тысячи). Товаром такие строки быть не
    могут — в товарной строке есть название. Не считаем их
    «не разобравшимися», чтобы не пугать людей цифрой."""
    vals = [to_text(c).lower().strip(" .,") for c in row if to_text(c)]
    return bool(vals) and all(v in ("0", "0.0", "итого", "итог") for v in vals)


def parse_rows(rows, cols, category, title, pet_per_liter=False):
    """Разбирает строки таблицы ПОСЛЕ шапки — общий конвейер для Excel и PDF.
    Возвращает (список позиций, число пропущенных строк)."""
    # «Двойной» режим — прайс, где в ячейках по ДВЕ цены (Black Cat):
    # строки разбирает parse_dual_row, а разделы прайса («Базовая
    # линейка», «Скидки») — это категории, не пивоварни.
    dual_mode = "price_dual" in cols or "price_dual_can" in cols

    items = []
    skipped = 0
    current_brewery = ""    # пивоварня из последней строки-заголовка раздела
    current_category = category   # категория (в двойном режиме меняется по разделам)
    for row in rows:
        # Строка-заголовок раздела: текст есть только в ОДНОЙ ячейке,
        # остальные пустые («Bakunin», «Achel Бельгия»).
        filled = [(idx, to_text(c)) for idx, c in enumerate(row) if to_text(c)]
        if len(filled) == 1 and filled[0][0] in (0, cols.get("name")):
            if dual_mode:
                # Раздел прайса («Скидки», «Постоянная линейка») — категория
                current_category = filled[0][1]
            else:
                # Обычный прайс: заголовок раздела — пивоварня
                current_brewery = clean_brewery(filled[0][1])
            continue

        if dual_mode:
            row_items = parse_dual_row(row, cols)
            if not row_items:
                # строка непустая, но позиций не дала (нет цен, «нет
                # в наличии») — считаем её пропущенной, если это не
                # нулевая заготовка формы заказа
                if any(to_text(c) for c in row) and not _is_junk_row(row):
                    skipped += 1
                continue
        else:
            item = parse_row(row, cols, pet_per_liter)
            if item is None:
                # строка непустая, но не разобралась — считаем её
                # пропущенной (кроме нулевых заготовок)
                if any(to_text(c) for c in row) and not _is_junk_row(row):
                    skipped += 1
                continue
            row_items = [item]

        for item in row_items:
            if not item["brewery"]:
                # в прайсе нет колонки «Пивоварня» — берём из заголовка раздела
                item["brewery"] = current_brewery
            item["sheet"] = title
            item["category"] = current_category
        items.extend(row_items)
    return items, skipped


def parse_pdf_file(path):
    """Читает PDF-прайс: pdfplumber находит на страницах таблицы,
    а дальше работает тот же конвейер, что и для Excel."""
    path = Path(path)
    supplier = guess_supplier(path.name)
    pet_per_liter = is_pet_per_liter_supplier(supplier)   # особенность поставщика

    items = []
    skipped = 0
    sheets_info = []

    last_cols = None    # раскладка колонок с прошлой страницы (память)
    with pdfplumber.open(path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            title = f"стр. {page_num}"       # в PDF нет названий листов — только страницы
            for table in page.extract_tables() or []:
                # Ячейки PDF — текст или «ничего»; пустоту превращаем в ""
                rows = [[to_text(c) for c in row] for row in table]
                header_row, cols = find_header_in_rows(rows)
                if header_row is not None:
                    # Строки НАД шапкой — хвост таблицы с прошлой страницы
                    # (так PDF рисует страницы: шапка нового раздела
                    # в середине). Читаем их по памяти, если она есть.
                    if last_cols is not None and header_row > 1:
                        pre_items, pre_skipped = parse_rows(
                            rows[:header_row - 1], last_cols, "", title,
                            pet_per_liter,
                        )
                        items.extend(pre_items)
                        skipped += pre_skipped
                        if pre_items:
                            sheets_info.append((title, len(pre_items)))
                    # Нашли шапку — запоминаем раскладку колонок
                    last_cols = cols
                    data_rows = rows[header_row:]
                elif last_cols is not None:
                    # Шапки нет: страница-продолжение таблицы с прошлой страницы.
                    # Читаем теми же колонками, что запомнили.
                    cols = last_cols
                    data_rows = rows
                else:
                    continue    # ни своей шапки, ни памяти — пропускаем
                # Категорию по PDF не угадать — оставляем пустой (НЕ «прочее»:
                # та категория не пускает позиции в топы, а тут кеги и банки вперемешку)
                page_items, page_skipped = parse_rows(
                    data_rows, cols, "", title, pet_per_liter
                )
                items.extend(page_items)
                skipped += page_skipped
                if page_items:
                    sheets_info.append((title, len(page_items)))

    return {
        "supplier": supplier,
        "items": items,
        "skipped": skipped,
        "sheets": sheets_info,
    }


def parse_price_file(path):
    """Читает прайс — Excel или PDF. Возвращает позиции + статистику."""
    path = Path(path)
    if path.suffix.lower() == ".pdf":
        return parse_pdf_file(path)

    supplier = guess_supplier(path.name)
    pet_per_liter = is_pet_per_liter_supplier(supplier)   # особенность поставщика
    wb = load_workbook(path, read_only=True, data_only=True)

    items = []
    skipped = 0
    sheets_info = []

    for ws in wb.worksheets:
        if ws.title.strip().lower() in SKIP_SHEETS:
            continue
        header_row, cols = find_header(ws)
        if header_row is None:
            continue    # не нашли шапку — лист не похож на таблицу с товарами

        title_lower = ws.title.lower()
        if "кег" in title_lower or "розлив" in title_lower:
            category = "кег"
        elif "фасов" in title_lower:
            category = "фасовка"
        elif "безалк" in title_lower:
            category = "безалкогольное"
        elif "прочее" in title_lower:
            # Лист честно называется «Прочее» (у WBT — лимонады Bench,
            # у Бирдилера — бокалы и оборудование): в топы не пускаем.
            category = "прочее"
        else:
            # Незнакомое название листа («Россия - крафт», «Импорт»,
            # «Акция») — это НЕ «прочее»! Не-пиво отсеют слова-пароли.
            category = ""

        sheet_items, sheet_skipped = parse_rows(
            ws.iter_rows(min_row=header_row + 1, values_only=True),
            cols, category, ws.title, pet_per_liter,
        )
        items.extend(sheet_items)
        skipped += sheet_skipped
        if sheet_items:
            sheets_info.append((ws.title, len(sheet_items)))

    wb.close()
    return {
        "supplier": supplier,
        "items": items,
        "skipped": skipped,
        "sheets": sheets_info,
    }


def build_summary(result, top=3) -> str:
    """Короткая сводка для чата: дешёвые кеги ЗА ЛИТР, банки ЗА ШТУКУ."""
    items = result["items"]

    # В топы берём только пиво, сидр и мёд — без газировок, соков и кваса.
    # (В данных и поиске они остаются! Убираем лишь из сводки.)
    # Проверяем слова в названии, стиле, заголовке раздела (пивоварня)
    # и НАЗВАНИИ ЛИСТА: у Price_Craft лимонады лежат на листе «Квас. Лимонад»,
    # а в PDF лимонады помечены только заголовком «Лимонады для HoReCa».
    def is_soft_drink(item):
        text = (
            item["name"] + " " + item["style"] + " " + item.get("brewery", "")
            + " " + item.get("category", "") + " " + item.get("sheet", "")
        ).lower()
        if any(word in text for word in SOFT_WORDS):
            return True
        return item.get("category") == "прочее"

    def short(text, limit=35):
        """Для сводки в чате: длинный стиль (в PDF в него слипается описание)
        обрезаем — целиком он остаётся в данных истории.
        Заодно перенос строк внутри текста меняем на пробел —
        иначе он ломает строчку в чате."""
        text = re.sub(r"\s+", " ", text or "").strip()
        return text if len(text) <= limit else text[:limit].rstrip() + "…"

    kegs = sorted(
        (i for i in items if i["keg"] and i["price_per_liter"] and not is_soft_drink(i)),
        key=lambda i: i["price_per_liter"],
    )
    cans = sorted(
        (i for i in items if not i["keg"] and i["price_per_liter"] and not is_soft_drink(i)),
        key=lambda i: i["price"],
    )

    # Объём (или цена упаковки) бывает не заполнен — например, в простой
    # Google-таблице без колонки объёма. Тогда просто НЕ пишем «за N л»
    # или «N ₽ =», а не падаем: позиция честно попадает в топ без деталей.
    def vol_text(item):
        return f" за {item['liters']:g} л" if item.get("liters") else ""

    def pack_price_text(item):
        return f"{item['price']:.0f} ₽{vol_text(item)} = " if item.get("price") else ""

    lines = []
    if kegs:
        lines.append("🏆 Самые дешёвые кеги и пэты (за литр):")
        for i in kegs[:top]:
            style = short(i["style"] or "стиль не указан")
            lines.append(
                f"  • {short(i['name'], 60)} ({style}) — "
                f"{pack_price_text(i)}"
                f"{i['price_per_liter']:.2f} ₽/л"
            )
    if cans:
        lines.append("🥫 Самые дешёвые банки и бутылки (за штуку):")
        for i in cans[:top]:
            style = short(i["style"] or "стиль не указан")
            lines.append(
                f"  • {short(i['name'], 60)} ({style}) — "
                f"{i['price']:.0f} ₽{vol_text(i)}"
            )
    return "\n".join(lines)


def save_history(result, source_path, history_dir):
    """Сохраняет разобранные позиции в CSV-файл истории (по датам).
    Позже это позволит команде /изменения сравнивать прайсы по дням."""
    history_dir = Path(history_dir)
    history_dir.mkdir(exist_ok=True)

    today = date.today().strftime("%Y-%m-%d")
    supplier = result["supplier"]
    csv_path = history_dir / f"{today} — {supplier} — данные.csv"

    fields = [
        "поставщик", "дата", "лист", "категория", "пивоварня", "название",
        "стиль", "тара", "объём штуки, л", "большая тара (кег/пэт)",
        "цена, руб", "цена за литр, руб",
    ]
    # utf-8-sig — чтобы файл сразу открывался в Excel без каши из букв
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for it in result["items"]:
            writer.writerow({
                "поставщик": result["supplier"],
                "дата": today,
                "лист": it.get("sheet", ""),
                "категория": it.get("category", ""),
                "пивоварня": it["brewery"],
                "название": it["name"],
                "стиль": it["style"],
                "тара": it["unit"],
                "объём штуки, л": it["liters"],
                "большая тара (кег/пэт)": "да" if it["keg"] else "нет",
                "цена, руб": it["price"],
                "цена за литр, руб": it["price_per_liter"],
            })
    return csv_path


def rename_last_supplier(history_dir, new_name: str):
    """Команда «назвать»: исправляет имя поставщика у ПОСЛЕДНЕГО
    сохранённого прайса — в данных истории и в имени CSV-файла.
    Присланный файл-оригинал в папке прайсы/ НЕ трогаем (правило:
    оригиналы остаются как лежат).

    Возвращает словарь:
      ok=True  — {старое, новое, файл}
      ok=False — {причина} — история пуста, имя пустое или у нового
                 поставщика уже есть прайс за ту же дату (не затираем!)."""
    history_dir = Path(history_dir)

    # Новое имя: обычные буквы (NFC), без «запрещённых» в Windows символов
    name = unicodedata.normalize("NFC", new_name or "").strip()
    name = re.sub(r'[\\/:*?"<>|]', " ", name)
    name = re.sub(r"\s+", " ", name).strip()
    if not name:
        return {"ok": False, "причина": "Пустое имя — напишите «назвать ИмяПоставщика»."}
    if len(name) > 60:
        return {"ok": False, "причина": "Имя слишком длинное (больше 60 символов)."}

    # Последний сохранённый прайс = CSV с самой поздней датой сохранения
    files = sorted(history_dir.glob("*.csv"), key=lambda p: p.stat().st_mtime)
    if not files:
        return {"ok": False, "причина": "История пуста — исправлять нечего."}
    old_path = files[-1]

    with open(old_path, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        fields = reader.fieldnames
        rows = list(reader)
    if not rows or not fields or "поставщик" not in fields:
        return {"ok": False, "причина": "Последний файл истории пустой или необычный."}

    old_name = rows[0]["поставщик"]
    row_date = rows[0]["дата"]

    new_path = history_dir / f"{row_date} — {name} — данные.csv"
    if new_path.exists() and new_path != old_path:
        return {
            "ok": False,
            "причина": (
                f"У поставщика «{name}» уже есть прайс за {row_date} — "
                "не буду затирать его. Скажите владельцу, он разберётся вручную."
            ),
        }

    # Меняем поставщика в каждой строке
    for r in rows:
        r["поставщик"] = name

    # Записываем обновлённый файл, старый убираем (это и есть «переименование»)
    with open(new_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    if old_path != new_path:
        old_path.unlink()

    return {"ok": True, "старое": old_name, "новое": name, "файл": new_path.name}